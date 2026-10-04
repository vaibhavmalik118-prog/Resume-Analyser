"""Resume Analyser - FastAPI backend that scores resumes with Google Gemini."""

import io
import json
import os
from datetime import date
from pathlib import Path

from docx import Document
from dotenv import load_dotenv
from fastapi import FastAPI, File, Form, HTTPException, UploadFile
from fastapi.responses import FileResponse
from fastapi.staticfiles import StaticFiles
from google import genai
from google.genai import errors, types

load_dotenv()

API_KEY = os.getenv("GEMINI_API_KEY", "").replace("your-gemini-api-key-here", "")
MODEL = os.getenv("GEMINI_MODEL", "gemini-3.8-flash")
FALLBACK_MODELS = [m.strip() for m in os.getenv("GEMINI_FALLBACK_MODELS", "gemini-3.7-flash,gemini-3.5-flash").split(",") if m.strip()]
MODELS = [MODEL] + [m for m in FALLBACK_MODELS if m != MODEL]
MAX_FILE_MB = int(os.getenv("MAX_FILE_MB", "10"))
SHORTLIST_THRESHOLD = int(os.getenv("SHORTLIST_THRESHOLD", "70"))  # job-match % needed to shortlist
FRONTEND_DIR = Path(__file__).parent / "frontend"

# Retry overloaded-server errors with exponential backoff (429 quota errors fall back to another model instead).
RETRY = types.HttpRetryOptions(attempts=4, initial_delay=2, max_delay=20,
                               http_status_codes=[500, 502, 503, 504])
client = genai.Client(api_key=API_KEY, http_options=types.HttpOptions(retry_options=RETRY)) if API_KEY else None
app = FastAPI(title="Resume Analyser")
app.mount("/static", StaticFiles(directory=FRONTEND_DIR), name="static")

SYSTEM_PROMPT = """You are an experienced technical recruiter and resume coach.
Evaluate the resume honestly and specifically: cite concrete details from the resume
rather than generic advice. Scores are 0-100, where 70 means a solid, competitive resume
and 90+ is reserved for exceptional ones. If a job description is provided, judge fit
against it; otherwise set job_match to null. If a candidate field is not in the document, use an
empty string (or 0 for years_experience) - never placeholders like "Not Provided" or "N/A".

Dates: today's date is given with each request. Treat "Present", "Current" or "Now" in a
date range as today's date. Compute years_experience by adding up the actual durations of
professional roles up to today (exclude education; count overlapping roles once; count
internships at most as partial experience), rounded to one decimal. Do not flag dates up to
today as being in the future."""

PLACEHOLDERS = {"not provided", "n/a", "na", "none", "unknown", "not specified", "not available", "-"}

STRING_LIST = {"type": "array", "items": {"type": "string"}}

ANALYSIS_SCHEMA = {
    "type": "object",
    "properties": {
        "candidate": {
            "type": "object",
            "properties": {
                "name": {"type": "string"},
                "email": {"type": "string"},
                "phone": {"type": "string"},
                "location": {"type": "string"},
                "current_title": {"type": "string"},
                "years_experience": {"type": "number"},
            },
            "required": ["name", "email", "phone", "location", "current_title", "years_experience"],
            "additionalProperties": False,
        },
        "overall_score": {"type": "integer"},
        "ats_score": {"type": "integer"},
        "summary": {"type": "string"},
        "section_scores": {
            "type": "array",
            "items": {
                "type": "object",
                "properties": {
                    "section": {"type": "string"},
                    "score": {"type": "integer"},
                    "feedback": {"type": "string"},
                },
                "required": ["section", "score", "feedback"],
                "additionalProperties": False,
            },
        },
        "skills": {
            "type": "object",
            "properties": {"technical": STRING_LIST, "soft": STRING_LIST},
            "required": ["technical", "soft"],
            "additionalProperties": False,
        },
        "strengths": STRING_LIST,
        "weaknesses": STRING_LIST,
        "suggestions": {
            "type": "array",
            "items": {
                "type": "object",
                "properties": {
                    "priority": {"type": "string", "enum": ["high", "medium", "low"]},
                    "area": {"type": "string"},
                    "suggestion": {"type": "string"},
                },
                "required": ["priority", "area", "suggestion"],
                "additionalProperties": False,
            },
        },
        "job_match": {
            "anyOf": [
                {
                    "type": "object",
                    "properties": {
                        "match_score": {"type": "integer"},
                        "matched_skills": STRING_LIST,
                        "missing_skills": STRING_LIST,
                        "verdict": {"type": "string"},
                    },
                    "required": ["match_score", "matched_skills", "missing_skills", "verdict"],
                    "additionalProperties": False,
                },
                {"type": "null"},
            ]
        },
    },
    "required": [
        "candidate", "overall_score", "ats_score", "summary", "section_scores",
        "skills", "strengths", "weaknesses", "suggestions", "job_match",
    ],
    "additionalProperties": False,
}


def resume_part(filename: str, data: bytes) -> types.Part:
    """Turn an uploaded file into a Gemini Part (PDFs are sent natively)."""
    ext = Path(filename).suffix.lower()
    if ext == ".pdf":
        return types.Part.from_bytes(data=data, mime_type="application/pdf")
    if ext == ".docx":
        doc = Document(io.BytesIO(data))
        text = "\n".join(p.text for p in doc.paragraphs)
        for table in doc.tables:
            for row in table.rows:
                text += "\n" + " | ".join(cell.text for cell in row.cells)
    elif ext in (".txt", ".md"):
        text = data.decode("utf-8", errors="replace")
    else:
        raise HTTPException(400, "Unsupported file type. Upload a PDF, DOCX or TXT file.")
    if not text.strip():
        raise HTTPException(400, "Could not find any text in that file.")
    return types.Part.from_text(text=f"<resume>\n{text}\n</resume>")


@app.get("/")
def index():
    return FileResponse(FRONTEND_DIR / "index.html")


@app.get("/api/health")
def health():
    return {"status": "ok", "models": MODELS, "api_key_configured": bool(API_KEY)}


@app.post("/api/analyze")
async def analyze(resume: UploadFile = File(...), job_description: str = Form("")):
    data = await resume.read()
    if not data:
        raise HTTPException(400, "The uploaded file is empty.")
    if len(data) > MAX_FILE_MB * 1024 * 1024:
        raise HTTPException(413, f"File is larger than {MAX_FILE_MB} MB.")

    instructions = f"Today's date is {date.today():%d %B %Y}. Analyse this resume."
    if job_description.strip():
        instructions += f"\n\nCompare it against this job description:\n<job_description>\n{job_description.strip()}\n</job_description>"

    if client is None:
        raise HTTPException(500, "GEMINI_API_KEY is not set - add it to your .env file and restart.")

    contents = [resume_part(resume.filename or "", data), instructions]
    config = types.GenerateContentConfig(
        system_instruction=SYSTEM_PROMPT,
        response_mime_type="application/json",
        response_json_schema=ANALYSIS_SCHEMA,
        temperature=0.2,
    )

    # Each model has its own free-tier quota and capacity, so on a 429 (quota) or a
    # persistent 5xx (overloaded) fall through to the next model.
    response = None
    last_server_error = None
    for model in MODELS:
        try:
            response = await client.aio.models.generate_content(model=model, contents=contents, config=config)
            break
        except errors.ClientError as e:
            if e.code == 429:
                print(f"[quota] {model} exhausted, trying next model")
                continue
            if e.code in (400, 401, 403) and "API key" in str(e):
                raise HTTPException(500, "Invalid GEMINI_API_KEY - check your .env file.")
            raise HTTPException(400, f"Gemini rejected the request: {e.message}")
        except errors.ServerError as e:
            # Still overloaded after the SDK's retries - another model may have capacity.
            print(f"[server] {model} returned {e.code}, trying next model")
            last_server_error = e.code
            continue
    if response is None:
        if last_server_error:
            raise HTTPException(502, f"Gemini servers are busy ({last_server_error}). Try again in a minute.")
        raise HTTPException(429, "The Gemini free-tier quota is used up for every configured model. "
                                 "It resets daily - try again later, or enable billing on your API key.")

    if not response.text:
        reason = response.candidates[0].finish_reason if response.candidates else "blocked"
        raise HTTPException(422, f"Gemini returned no analysis (reason: {reason}).")
    try:
        result = json.loads(response.text)
    except json.JSONDecodeError:
        raise HTTPException(502, "Gemini returned an unreadable response. Please try again.")

    # Blank out placeholder values so the UI never shows "Not Provided" as if it were real data.
    for key, value in result.get("candidate", {}).items():
        if isinstance(value, str) and value.strip().lower() in PLACEHOLDERS:
            result["candidate"][key] = ""

    # Shortlisting decision is a fixed rule on our side, not left to the model.
    if result.get("job_match"):
        result["job_match"]["threshold"] = SHORTLIST_THRESHOLD
        result["job_match"]["shortlisted"] = result["job_match"]["match_score"] >= SHORTLIST_THRESHOLD
    return result


if __name__ == "__main__":
    import uvicorn

    uvicorn.run("main:app", host="127.0.0.1", port=int(os.getenv("PORT", "8000")), reload=True)
